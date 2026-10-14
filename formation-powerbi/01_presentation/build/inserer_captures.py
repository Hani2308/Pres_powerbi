"""Insère les captures PNG de 05_captures/ dans la présentation, à la place des emplacements gris.

Pour chaque forme nommée « Capture Cxx » du deck :
- si 05_captures/Cxx.png existe, l'image est posée dans le cadre de l'emplacement (proportions
  conservées, centrée), placée sous les annotations rouges, avec un texte de remplacement ;
- l'emplacement gris et sa légende « Capture Cxx » sont supprimés.
Le nom du formateur est aussi remplacé par celui passé en option.

Usage : python inserer_captures.py [--nom "SAHIR Hani"] [--captures dossier] [--sortie fichier.pptx]
Sortie : 01_presentation/Initiation_Power_BI_complete.pptx (la version avec emplacements reste intacte).
"""
import argparse
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent.parent
SOURCE = ICI.parent / "Initiation_Power_BI.pptx"
SORTIE = ICI.parent / "Initiation_Power_BI_complete.pptx"
DOSSIER = RACINE / "05_captures"


def remplacer_nom(prs, nom):
    """Remplace « Hani Sahir » dans les masques, les dispositions et les slides."""
    def dans(shapes):
        for sh in shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        if "Hani Sahir" in run.text:
                            run.text = run.text.replace("Hani Sahir", nom)
    for master in prs.slide_masters:
        dans(master.shapes)
        for layout in master.slide_layouts:
            dans(layout.shapes)
    for slide in prs.slides:
        dans(slide.shapes)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--nom", default="SAHIR Hani")
    ap.add_argument("--captures", default=str(DOSSIER), help="dossier des PNG (défaut : 05_captures)")
    ap.add_argument("--sortie", default=str(SORTIE))
    args = ap.parse_args()
    dossier = Path(args.captures)

    descriptions = {c["id"]: c["quoi"] for c in json.loads((ICI / "captures.json").read_text(encoding="utf-8"))}
    prs = Presentation(SOURCE)
    inserees, manquantes = [], []

    for n, slide in enumerate(prs.slides, 1):
        emplacements = [sh for sh in slide.shapes if sh.name.startswith("Capture C")]
        for emp in emplacements:
            cid = emp.name.split(" ", 1)[1]
            png = dossier / f"{cid}.png"
            if not png.exists():
                manquantes.append(f"{cid} (slide {n})")
                continue
            with Image.open(png) as im:
                iw, ih = im.size
            echelle = min(emp.width / iw, emp.height / ih)
            w, h = int(iw * echelle), int(ih * echelle)
            x = emp.left + (emp.width - w) // 2
            y = emp.top + (emp.height - h) // 2
            pic = slide.shapes.add_picture(str(png), x, y, w, h)
            pic.name = f"Image {cid}"
            pic._element.nvPicPr.cNvPr.set("descr", descriptions.get(cid, cid))
            # Même position dans l'empilement que l'emplacement : les annotations restent au-dessus.
            emp._element.addprevious(pic._element)
            for sh in list(slide.shapes):
                if sh.name in (f"Capture {cid}", f"Légende {cid}"):
                    sh._element.getparent().remove(sh._element)
            inserees.append(cid)

    remplacer_nom(prs, args.nom)
    prs.save(args.sortie)
    print(f"{len(inserees)} captures insérées -> {args.sortie}")
    if manquantes:
        print(f"{len(manquantes)} captures manquantes, emplacement laissé en place : {', '.join(manquantes)}")
        sys.exit(1)


if __name__ == "__main__":
    main()
