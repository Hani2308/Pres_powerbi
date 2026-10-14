"""Génère le fil rouge de la formation : l'export de caisse d'Atlas Bazar (enseigne fictive).

Produit, dans le dossier de ce script :
  Ventes_Retail.xlsx          export au 30/09/2026 + feuille « Corrigé formateur »
  Ventes_Retail_MAJ.xlsx      même export, arrêté au 11/10/2026 (démo de l'actualisation)
  participants/               les deux fichiers sans la feuille de corrigé
  valeurs_attendues.json      résultats des 7 formules DAX, recalculés ici en pandas

Graine fixe : relancer le script redonne exactement les mêmes fichiers.
Usage : python generer_dataset.py
"""

import json
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ICI = Path(__file__).resolve().parent
GRAINE = 59

DEBUT = date(2024, 10, 1)
FIN = date(2026, 9, 30)
FIN_MAJ = date(2026, 10, 11)
BASE_TICKETS = 0.27  # calibré pour ~5 000 lignes sur la période principale

# Code, nom, « Ville - Région » tel que saisi dans le référentiel, surface, ouverture, poids commercial
MAGASINS = [
    ("MAG-01", "Atlas Bazar Maârif", "Casablanca - Casablanca-Settat", 420, date(2014, 3, 1), 1.5),
    ("MAG-02", "Atlas Bazar Aïn Sebaâ", "casa - Casablanca-Settat", 310, date(2016, 9, 15), 1.2),
    ("MAG-03", "Atlas Bazar Hay Hassani", "CASABLANCA  - Casablanca-Settat", 280, date(2018, 5, 2), 1.1),
    ("MAG-04", "Atlas Bazar Agdal", "Rabat - Rabat-Salé-Kénitra", 350, date(2015, 1, 10), 1.3),
    ("MAG-05", "Atlas Bazar Tabriquet", "Salé - Rabat-Salé-Kénitra", 240, date(2019, 11, 4), 0.9),
    ("MAG-06", "Atlas Bazar Guéliz", "Marrakech - Marrakech-Safi", 330, date(2016, 4, 20), 1.2),
    ("MAG-07", "Atlas Bazar Atlas Fès", "Fès - Fès-Meknès", 300, date(2017, 6, 12), 1.0),
    ("MAG-08", "Atlas Bazar Hamria", "Meknès - Fès-Meknès", 220, date(2020, 2, 3), 0.8),
    ("MAG-09", "Atlas Bazar Iberia", "Tanger - Tanger-Tétouan-Al Hoceïma", 310, date(2018, 10, 8), 1.1),
    ("MAG-10", "Atlas Bazar Talborjt", "Agadir - Souss-Massa", 260, date(2019, 3, 18), 0.9),
    ("MAG-11", "Atlas Bazar Oujda Centre", "Oujda - Oriental", 200, date(2021, 9, 1), 0.7),
    ("MAG-12", "Atlas Bazar Tétouan Médina", "Tétouan - Tanger-Tétouan-Al Hoceïma", 210, date(2022, 4, 25), 0.8),
]
# Ouvert en mars 2026, vend, mais personne ne l'a ajouté au référentiel.
MAG13 = ("MAG-13", 0.9, date(2026, 3, 16))

CATEGORIES = {
    # catégorie : part de base des lignes de ticket
    "Épicerie": 0.36,
    "Boissons": 0.21,
    "Hygiène et beauté": 0.15,
    "Maison": 0.10,
    "Textile": 0.13,
    "Fournitures scolaires": 0.025,
}

PRODUITS = [
    # Réf, désignation, catégorie, prix catalogue MAD, quantité max par ligne
    ("EPI-001", "Huile d'olive vierge 1 L", "Épicerie", 75.0, 3),
    ("EPI-002", "Thé vert gunpowder 200 g", "Épicerie", 18.0, 4),
    ("EPI-003", "Dattes Medjool 1 kg", "Épicerie", 90.0, 3),
    ("EPI-004", "Farine de blé 5 kg", "Épicerie", 45.0, 2),
    ("EPI-005", "Sucre en pain 2 kg", "Épicerie", 26.0, 3),
    ("EPI-006", "Lentilles 1 kg", "Épicerie", 16.0, 3),
    ("EPI-007", "Pois chiches 1 kg", "Épicerie", 14.0, 3),
    ("EPI-008", "Couscous moyen 1 kg", "Épicerie", 15.0, 4),
    ("EPI-009", "Sardines à l'huile 125 g", "Épicerie", 9.0, 6),
    ("EPI-010", "Café moulu 250 g", "Épicerie", 38.0, 2),
    ("EPI-011", "Miel de thym 250 g", "Épicerie", 65.0, 2),
    ("EPI-012", "Amandes décortiquées 500 g", "Épicerie", 70.0, 2),
    ("EPI-013", "Chebakia 500 g", "Épicerie", 55.0, 3),
    ("EPI-014", "Vermicelles 500 g", "Épicerie", 8.5, 5),
    ("BOI-001", "Eau minérale 1,5 L", "Boissons", 6.0, 12),
    ("BOI-002", "Soda cola 1 L", "Boissons", 10.0, 6),
    ("BOI-003", "Jus d'orange 1 L", "Boissons", 14.0, 4),
    ("BOI-004", "Lait demi-écrémé 1 L", "Boissons", 7.5, 6),
    ("BOI-005", "Lben 50 cl", "Boissons", 4.0, 6),
    ("BOI-006", "Eau gazeuse 1 L", "Boissons", 7.0, 6),
    ("BOI-007", "Sirop de grenadine 75 cl", "Boissons", 22.0, 2),
    ("HYG-001", "Savon beldi 250 g", "Hygiène et beauté", 25.0, 3),
    ("HYG-002", "Shampoing 400 ml", "Hygiène et beauté", 32.0, 2),
    ("HYG-003", "Dentifrice 75 ml", "Hygiène et beauté", 18.0, 3),
    ("HYG-004", "Ghassoul 200 g", "Hygiène et beauté", 20.0, 2),
    ("HYG-005", "Déodorant 150 ml", "Hygiène et beauté", 35.0, 2),
    ("HYG-006", "Couches taille 4 x 44", "Hygiène et beauté", 120.0, 2),
    ("HYG-007", "Eau de fleur d'oranger 250 ml", "Hygiène et beauté", 16.0, 3),
    ("HYG-008", "Huile d'argan cosmétique 50 ml", "Hygiène et beauté", 85.0, 1),
    ("MAI-001", "Théière inox 1 L", "Maison", 180.0, 1),
    ("MAI-002", "Verres à thé x 6", "Maison", 60.0, 2),
    ("MAI-003", "Tajine terre cuite 30 cm", "Maison", 90.0, 1),
    ("MAI-004", "Couteau de boucher", "Maison", 45.0, 2),
    ("MAI-005", "Plateau à thé cuivre", "Maison", 110.0, 1),
    ("MAI-006", "Seau plastique 10 L", "Maison", 25.0, 2),
    ("MAI-007", "Nappe brodée", "Maison", 70.0, 1),
    ("MAI-008", "Kanoun (brasero) terre cuite", "Maison", 150.0, 1),
    ("MAI-009", "Lot de 3 couteaux", "Maison", 130.0, 1),
    ("MAI-010", "Grille barbecue", "Maison", 65.0, 1),
    ("TEX-001", "Djellaba enfant", "Textile", 250.0, 1),
    ("TEX-002", "Gandoura homme", "Textile", 220.0, 1),
    ("TEX-003", "T-shirt coton", "Textile", 60.0, 3),
    ("TEX-004", "Babouches cuir", "Textile", 140.0, 2),
    ("TEX-005", "Pyjama enfant", "Textile", 120.0, 2),
    ("TEX-006", "Foulard", "Textile", 50.0, 2),
    ("TEX-007", "Caftan femme", "Textile", 380.0, 1),
    ("TEX-008", "Chaussettes x 3", "Textile", 30.0, 3),
    ("FOU-001", "Cartable", "Fournitures scolaires", 180.0, 1),
    ("FOU-002", "Cahier 96 pages", "Fournitures scolaires", 8.0, 10),
    ("FOU-003", "Lot de 10 stylos", "Fournitures scolaires", 15.0, 3),
    ("FOU-004", "Trousse", "Fournitures scolaires", 35.0, 2),
    ("FOU-005", "Calculatrice scientifique", "Fournitures scolaires", 90.0, 1),
    ("FOU-006", "Tablier d'écolier", "Fournitures scolaires", 60.0, 2),
    ("FOU-007", "Protège-cahiers x 10", "Fournitures scolaires", 12.0, 3),
]

# Dates de fêtes : estimations (calendrier lunaire, annonce officielle à J-1). Voir le corrigé.
RAMADAN = [(date(2025, 3, 2), date(2025, 3, 30)), (date(2026, 2, 19), date(2026, 3, 19))]
AID_FITR = [date(2025, 3, 31), date(2026, 3, 20)]
# 2025 : appel royal à ne pas accomplir le sacrifice (sécheresse), d'où un effet faible.
AID_ADHA = [(date(2025, 6, 7), 1.3), (date(2026, 5, 27), 3.0)]

MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]
JOUR_SEMAINE = [0.9, 0.9, 0.95, 1.0, 0.85, 1.25, 1.2]  # lundi -> dimanche


def multiplicateurs(j: date) -> dict:
    """Effet saisonnier par catégorie pour le jour j."""
    m = {c: 1.0 for c in CATEGORIES}
    prix_textile = 1.0
    for debut, fin in RAMADAN:
        if debut <= j <= fin:
            m["Épicerie"] *= 3.0
            m["Boissons"] *= 2.0
            m["Maison"] *= 1.2
            m["Textile"] *= 0.8
            if (fin - j).days < 8:
                m["Textile"] *= 3.5
    if j in AID_FITR:
        m = {c: v * 0.4 for c, v in m.items()}
    for jour_aid, force in AID_ADHA:
        ecart = (jour_aid - j).days
        if 0 < ecart <= 10:
            m["Maison"] *= force
            m["Épicerie"] *= 1 + (force - 1) * 0.15
        if ecart == 0:
            m = {c: v * 0.4 for c, v in m.items()}
    if (j.month == 8 and j.day >= 25) or (j.month == 9 and j.day <= 15):
        m["Fournitures scolaires"] *= 10
        m["Textile"] *= 1.3
    if j.month == 1 or j.month in (7, 8):
        m["Textile"] *= 1.6
        prix_textile = 0.7
    if j.month in (6, 7, 8):
        m["Boissons"] *= 1.5
    return m, prix_textile


def generer_ventes(rng, fin: date) -> pd.DataFrame:
    produits = pd.DataFrame(PRODUITS, columns=["ref", "nom", "cat", "prix", "qmax"])
    par_cat = {c: produits[produits.cat == c].reset_index(drop=True) for c in CATEGORIES}
    magasins = [(m[0], m[5], None) for m in MAGASINS] + [MAG13]
    lignes = []
    j = DEBUT
    while j <= fin:
        mult, prix_textile = multiplicateurs(j)
        poids = {c: CATEGORIES[c] * mult[c] for c in CATEGORIES}
        total_poids = sum(poids.values())
        intensite = total_poids / sum(CATEGORIES.values())
        tendance = 1 + 0.07 * (j - DEBUT).days / 365
        cats = list(poids)
        p_cats = np.array([poids[c] for c in cats]) / total_poids
        for code, force, ouverture in magasins:
            if ouverture and j < ouverture:
                continue
            lam = BASE_TICKETS * force * JOUR_SEMAINE[j.weekday()] * tendance * intensite
            for n in range(1, rng.poisson(lam) + 1):
                ticket = f"T{code[-2:]}-{j:%y%m%d}-{n:03d}"
                nb_lignes = rng.choice([1, 2, 3], p=[0.5, 0.33, 0.17])
                deja = set()
                for _ in range(nb_lignes):
                    cat = cats[rng.choice(len(cats), p=p_cats)]
                    prod = par_cat[cat].iloc[rng.integers(len(par_cat[cat]))]
                    if prod.ref in deja:
                        continue
                    deja.add(prod.ref)
                    prix = prod.prix * (prix_textile if cat == "Textile" else 1.0)
                    if rng.random() < 0.04:  # promo ponctuelle en magasin
                        prix *= 0.9
                    lignes.append({
                        "Date": j,
                        "N° ticket": ticket,
                        "Code magasin": code,
                        "Réf produit": prod.ref,
                        "Quantité": int(rng.integers(1, prod.qmax + 1)),
                        "Prix unitaire": round(prix, 2),
                    })
        j += timedelta(days=1)
    return pd.DataFrame(lignes)


def calendrier() -> pd.DataFrame:
    jours = pd.date_range("2024-01-01", "2026-12-31", freq="D")
    return pd.DataFrame({
        "Date": jours.date,
        "Année": jours.year,
        "N° mois": jours.month,
        "Mois": [MOIS_FR[m - 1] for m in jours.month],
        "Trimestre": ["T" + str(q) for q in jours.quarter],
    })


BLOC_DOUBLON = (date(2026, 6, 8), date(2026, 6, 12))  # semaine réexportée et recollée


def valeurs_attendues(propre: pd.DataFrame, brut_lignes: int, doublons: int) -> dict:
    """Résultats que doivent trouver les participants, une fois le nettoyage fait."""
    v = propre.copy()
    v["Montant"] = v["Quantité"] * v["Prix unitaire"]
    v["Date"] = pd.to_datetime(v["Date"])
    ville = {m[0]: m[2].split(" - ")[0].strip().title().replace("Casa", "Casablanca")
             .replace("Casablancablanca", "Casablanca") for m in MAGASINS}
    v["Ville"] = v["Code magasin"].map(ville)
    ca = v["Montant"].sum()
    nb_tickets = v["N° ticket"].nunique()

    def ca_mois(a, m):
        return v[(v.Date.dt.year == a) & (v.Date.dt.month == m)]["Montant"].sum()

    par_mois = (v.groupby([v.Date.dt.year.rename("Année"), v.Date.dt.month.rename("Mois")])["Montant"]
                .sum().round(2))
    mensuel = [{"annee": int(a), "mois": int(m), "ca": float(x)} for (a, m), x in par_mois.items()]
    mars26, mars25 = ca_mois(2026, 3), ca_mois(2025, 3)
    fev26, fev25 = ca_mois(2026, 2), ca_mois(2025, 2)
    return {
        "lignes_brutes_export": brut_lignes,
        "lignes_dupliquees": doublons,
        "lignes_apres_nettoyage": int(len(v)),
        "ca_total": round(float(ca), 2),
        "ca_doublons_si_oublies": None,  # rempli par main()
        "nb_tickets": int(nb_tickets),
        "panier_moyen": round(float(ca / nb_tickets), 2),
        "ca_casablanca": round(float(v[v.Ville == "Casablanca"]["Montant"].sum()), 2),
        "ca_mag13_vide": round(float(v[v["Code magasin"] == "MAG-13"]["Montant"].sum()), 2),
        "ca_par_ville": {k: round(float(x), 2) for k, x in v.groupby("Ville")["Montant"].sum().items()},
        "ca_fevrier_2026": round(float(fev26), 2), "ca_fevrier_2025": round(float(fev25), 2),
        "evol_fevrier_2026": round(float(fev26 / fev25 - 1), 4),
        "ca_mars_2026": round(float(mars26), 2), "ca_mars_2025": round(float(mars25), 2),
        "evol_mars_2026": round(float(mars26 / mars25 - 1), 4),
        "ca_annee_2025": round(float(v[v.Date.dt.year == 2025]["Montant"].sum()), 2),
        "mensuel": mensuel,
    }


def ecrire_classeur(chemin: Path, ventes: pd.DataFrame, fin: date, corrige: dict | None):
    wb = Workbook()
    police = Font(name="Arial", size=10)
    gras = Font(name="Arial", size=10, bold=True)

    # Export_Ventes, tel que la caisse le sort
    ws = wb.active
    ws.title = "Export_Ventes"
    genere = fin + timedelta(days=1)
    ws["A1"] = (f"ATLAS BAZAR - Journal des ventes - Période du {DEBUT:%d/%m/%Y} au {fin:%d/%m/%Y} "
                f"- Généré le {genere:%d/%m/%Y} à 07:42")
    ws["A1"].font = gras
    entetes = ["Date", "N° ticket", "Code magasin", "Réf produit", "Quantité", "Prix unitaire"]
    for c, h in enumerate(entetes, 1):
        ws.cell(row=3, column=c, value=h).font = gras
    doublons = ventes[(ventes.Date >= BLOC_DOUBLON[0]) & (ventes.Date <= BLOC_DOUBLON[1])].copy()
    doublons["Code magasin"] = doublons["Code magasin"] + " "
    tout = pd.concat([ventes, doublons], ignore_index=True)
    r = 4
    for row in tout.itertuples(index=False):
        ws.cell(row=r, column=1, value=row[0].strftime("%d/%m/%Y")).number_format = "@"
        for c, val in enumerate(row[1:], 2):
            ws.cell(row=r, column=c, value=val)
        ws.cell(row=r, column=6).number_format = "0.00"
        r += 1
    premiere_dup = 4 + len(ventes)
    derniere_donnee = r - 1
    ligne_total = r + 1  # une ligne vide, puis le total
    ws.cell(row=ligne_total, column=1, value="TOTAL GÉNÉRAL").font = gras
    ws.cell(row=ligne_total, column=5, value=int(tout["Quantité"].sum())).font = gras
    for row in ws.iter_rows(min_row=2, max_row=ligne_total):
        for cell in row:
            if not cell.font.bold:
                cell.font = police
    for c, w in zip("ABCDEF", [12, 18, 14, 12, 10, 13]):
        ws.column_dimensions[c].width = w

    # Magasins : référentiel tenu à la main
    wm = wb.create_sheet("Magasins")
    for c, h in enumerate(["Code magasin", "Nom magasin", "Ville - Région", "Surface (m²)", "Date d'ouverture"], 1):
        wm.cell(row=1, column=c, value=h).font = gras
    lignes_mag = [m[:5] for m in MAGASINS]
    lignes_mag.insert(5, ("MAG-05", "Atlas Bazar Tabriquet", "Salé  - Rabat-Salé-Kénitra", 240, date(2019, 11, 4)))
    for i, m in enumerate(lignes_mag, 2):
        for c, val in enumerate(m, 1):
            cell = wm.cell(row=i, column=c, value=val)
            cell.font = police
            if c == 5:
                cell.number_format = "DD/MM/YYYY"
    for c, w in zip("ABCDE", [14, 28, 38, 13, 16]):
        wm.column_dimensions[c].width = w

    wp = wb.create_sheet("Produits")
    for c, h in enumerate(["Réf produit", "Désignation", "Catégorie", "Prix catalogue (MAD)"], 1):
        wp.cell(row=1, column=c, value=h).font = gras
    for i, p in enumerate(PRODUITS, 2):
        for c, val in enumerate(p[:4], 1):
            wp.cell(row=i, column=c, value=val).font = police
        wp.cell(row=i, column=4).number_format = "0.00"
    for c, w in zip("ABCD", [12, 32, 22, 20]):
        wp.column_dimensions[c].width = w

    wc = wb.create_sheet("Calendrier")
    cal = calendrier()
    for c, h in enumerate(cal.columns, 1):
        wc.cell(row=1, column=c, value=h).font = gras
    for i, row in enumerate(cal.itertuples(index=False), 2):
        for c, val in enumerate(row, 1):
            cell = wc.cell(row=i, column=c, value=val)
            cell.font = police
            if c == 1:
                cell.number_format = "DD/MM/YYYY"
    for c, w in zip("ABCDE", [12, 8, 9, 12, 10]):
        wc.column_dimensions[c].width = w

    if corrige is not None:
        ecrire_corrige(wb, corrige, premiere_dup, derniere_donnee, ligne_total, len(doublons))
    wb.save(chemin)
    return len(tout), len(doublons), premiere_dup, derniere_donnee, ligne_total


def ecrire_corrige(wb, v, premiere_dup, derniere_donnee, ligne_total, nb_dup):
    ws = wb.create_sheet("Corrigé formateur")
    titre = Font(name="Arial", size=12, bold=True)
    gras = Font(name="Arial", size=10, bold=True)
    police = Font(name="Arial", size=10)
    fond = PatternFill("solid", fgColor="E8E4DA")
    enveloppe = Alignment(wrap_text=True, vertical="top")
    ws["A1"] = "Corrigé formateur : défauts plantés dans l'export et valeurs de contrôle"
    ws["A1"].font = titre
    ws["A2"] = ("Feuille réservée au formateur. Les participants reçoivent le dossier participants/, "
                "où elle n'existe pas.")
    ws["A2"].font = police

    entetes = ["N°", "Feuille", "Défaut planté", "Où exactement", "Ce qu'on voit dans Power BI si on l'oublie",
               "Correction dans Power Query", "Slide"]
    defauts = [
        ("Export_Ventes", "Deux lignes de titre au-dessus des en-têtes",
         "Lignes 1 et 2 (A1 : titre de l'export, ligne 2 vide). En-têtes en ligne 3.",
         "Colonnes nommées Column1, Column2… ; la première ligne de données contient le titre.",
         "Accueil > Supprimer les lignes > Supprimer les lignes du haut (2), puis Utiliser la première ligne pour les en-têtes.",
         "9, 16"),
        ("Export_Ventes", "Dates en texte au format JJ/MM/AAAA",
         f"Colonne A, lignes 4 à {derniere_donnee} (format de cellule Texte).",
         "Sur un poste ou un fichier réglé en Anglais (États-Unis) : « Error » pour tout jour > 12, "
         "jour et mois inversés sans message pour les jours ≤ 12.",
         "Icône de type de la colonne Date > Utilisation des paramètres régionaux… > Date, Français (Maroc).",
         "14, 15"),
        ("Export_Ventes", "Ligne vide puis ligne « TOTAL GÉNÉRAL » collées sous les données",
         f"Ligne {ligne_total - 1} vide ; ligne {ligne_total} : A = « TOTAL GÉNÉRAL », E = total des quantités.",
         "Avec Charger direct : la somme des quantités est doublée ; une erreur de chargement sur la colonne Date.",
         "Filtrer la colonne N° ticket pour exclure (null) : enlève les deux lignes, même si l'export change de "
         "longueur. Alternative fragile : Supprimer les lignes du bas (2).",
         "10, 16"),
        ("Export_Ventes", "Semaine réexportée et recollée à la fin, avec un espace après chaque code magasin",
         f"Lignes {premiere_dup} à {derniere_donnee} ({nb_dup} lignes) : copies exactes des ventes du "
         f"{BLOC_DOUBLON[0]:%d/%m/%Y} au {BLOC_DOUBLON[1]:%d/%m/%Y}, sauf « MAG-xx » devenu « MAG-xx␣ ».",
         f"CA surestimé de {v['ca_doublons_si_oublies']:,.2f} MAD. Supprimer les doublons AVANT de supprimer les "
         "espaces n'enlève aucune de ces lignes.".replace(",", " "),
         "Transformer > Format > Supprimer les espaces sur Code magasin, PUIS Accueil > Supprimer les lignes > "
         "Supprimer les doublons.",
         "20, 21"),
        ("Export_Ventes", "Magasin MAG-13 (Kénitra, ouvert le 16/03/2026) absent du référentiel",
         "Colonne C, toutes les lignes « MAG-13 » à partir du 16/03/2026.",
         f"Une ligne « (Vide) » dans tout visuel par ville ou par magasin : {v['ca_mag13_vide']:,.2f} MAD.".replace(",", " "),
         "Ne pas corriger dans Power BI : demander l'ajout de MAG-13 au propriétaire du référentiel. "
         "En salle, on le laisse pour montrer l'effet.",
         "27, 29"),
        ("Magasins", "Ville et région fusionnées dans une colonne « Ville - Région »",
         "Colonne C, lignes 2 à 14.",
         "Impossible de filtrer par ville ou par région.",
         "Fractionner la colonne > Par délimiteur, délimiteur personnalisé « espace-tiret-espace » ( - ). "
         "Un simple tiret couperait aussi « Casablanca-Settat ».",
         "18"),
        ("Magasins", "Casse incohérente : « Casablanca », « casa », « CASABLANCA »",
         "C2 (Casablanca), C3 (casa), C4 (CASABLANCA suivi de deux espaces).",
         "Power Query voit trois villes. Une fois chargé, le modèle fusionne « CASABLANCA » et « Casablanca » "
         "(il ignore la casse) mais « casa » reste une ville à part.",
         "Supprimer les espaces, Majuscule à chaque mot, puis Remplacer les valeurs « Casa » par « Casablanca » "
         "avec l'option « Mettre en correspondance tout le contenu de la cellule ». Sans cette case : « Casablancablanca ».",
         "17, 19"),
        ("Magasins", "Espaces en fin de ville",
         "C4 (« CASABLANCA  »), C7 (« Salé  » sur la ligne en double).",
         "Deux « Salé » distincts dans les listes de filtre de Power Query.",
         "Transformer > Format > Supprimer les espaces, après le fractionnement.",
         "19"),
        ("Magasins", "MAG-05 saisi deux fois",
         "Lignes 6 et 7 : même code, la seconde avec « Salé » suivi d'espaces.",
         "La relation Magasins - Ventes ne peut pas être un-à-plusieurs : Power BI propose plusieurs-à-plusieurs.",
         "Supprimer les espaces sur la ville, puis Supprimer les doublons. Vérifier : 12 lignes.",
         "26"),
        ("Calendrier", "Mois en toutes lettres (aucun défaut de saisie)",
         "Colonne D.",
         "Axe des mois trié par ordre alphabétique : août, avril, décembre, février…",
         "Pas dans Power Query : vue Table, colonne Mois > Outils de colonne > Trier par colonne > N° mois.",
         "37"),
    ]
    for c, h in enumerate(entetes, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = gras
        cell.fill = fond
        cell.alignment = enveloppe
    for i, d in enumerate(defauts, 5):
        ws.cell(row=i, column=1, value=i - 4).font = police
        for c, val in enumerate(d, 2):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = police
            cell.alignment = enveloppe

    r = 5 + len(defauts) + 2
    ws.cell(row=r, column=1, value="Valeurs de contrôle").font = titre
    r += 1
    for c, h in enumerate(["Indicateur", "Valeur attendue", "Comment elle est obtenue"], 1):
        cell = ws.cell(row=r, column=c if c == 1 else c + 1, value=h)
        cell.font = gras
        cell.fill = fond
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1
    controles = [
        ("Lignes de données dans l'export brut (hors titre, en-têtes, ligne vide et total)",
         f"=COUNTA(Export_Ventes!B4:B{derniere_donnee})", "Formule Excel sur la feuille Export_Ventes.", "0"),
        ("dont lignes recollées en double", f"=COUNTA(Export_Ventes!B{premiere_dup}:B{derniere_donnee})",
         "Formule Excel sur le bloc recollé.", "0"),
        ("Lignes de la table Ventes après nettoyage", v["lignes_apres_nettoyage"],
         "generer_dataset.py (pandas). À retrouver dans Power Query, en bas à gauche de l'aperçu.", "0"),
        ("CA total (MAD) : formule 2", v["ca_total"], "Somme de Quantité x Prix unitaire après nettoyage.", "#,##0.00"),
        ("Nb tickets : formule 3", v["nb_tickets"], "Nombre de N° ticket distincts.", "#,##0"),
        ("Panier moyen (MAD) : formule 4", v["panier_moyen"], "CA / Nb tickets.", "#,##0.00"),
        ("CA Casablanca (MAD) : formule 5", v["ca_casablanca"], "MAG-01 + MAG-02 + MAG-03.", "#,##0.00"),
        ("CA mars 2026 (MAD)", v["ca_mars_2026"], "Ramadan 2026 : 19/02 au 19/03 (estimation).", "#,##0.00"),
        ("CA N-1 de mars 2026 : formule 6", v["ca_mars_2025"], "Ramadan 2025 : 02/03 au 30/03 (estimation).", "#,##0.00"),
        ("Évolution vs N-1, mars 2026 : formule 7", v["evol_mars_2026"], "Baisse due au décalage du Ramadan.", "0.0%"),
        ("Évolution vs N-1, février 2026", v["evol_fevrier_2026"], "Hausse due au même décalage.", "0.0%"),
        ("CA de la ligne « (Vide) » (MAG-13)", v["ca_mag13_vide"], "Ventes de MAG-13 depuis le 16/03/2026.", "#,##0.00"),
    ]
    for lib, val, comment, fmt in controles:
        ws.cell(row=r, column=1, value=lib).font = police
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        cell = ws.cell(row=r, column=3, value=val)
        cell.font = police
        cell.number_format = fmt
        ws.cell(row=r, column=4, value=comment).font = police
        r += 1
    r += 1
    notes = [
        "Les valeurs sans formule ont été calculées par generer_dataset.py (graine 59) et sont recopiées ici ; "
        "elles servent de référence pour vérifier les mesures DAX en salle.",
        "Dates de fêtes : Ramadan, Aïd al-Fitr et Aïd al-Adha suivent le calendrier lunaire et sont fixées au Maroc "
        "la veille. Les dates utilisées sont des estimations, suffisantes pour un jeu de données de formation.",
        "Aïd al-Adha 2025 : appel royal à ne pas accomplir le sacrifice (sécheresse). Les ventes Maison montent "
        "beaucoup moins qu'en 2026. Anecdote utilisable si un participant le remarque.",
        "Les données sont un extrait fictif : une enseigne de 12 magasins vend bien plus. On le dit en salle si "
        "quelqu'un trouve le CA petit.",
    ]
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = police
        r += 1
    for c, w in zip("ABCDEFG", [5, 16, 34, 40, 44, 52, 8]):
        ws.column_dimensions[c].width = w
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["C"].width = 34


def main():
    rng = np.random.default_rng(GRAINE)
    ventes_maj = generer_ventes(rng, FIN_MAJ)
    ventes = ventes_maj[ventes_maj.Date <= FIN].reset_index(drop=True)

    doublons = ventes[(ventes.Date >= BLOC_DOUBLON[0]) & (ventes.Date <= BLOC_DOUBLON[1])]
    v = valeurs_attendues(ventes, len(ventes) + len(doublons), len(doublons))
    v["ca_doublons_si_oublies"] = round(float((doublons["Quantité"] * doublons["Prix unitaire"]).sum()), 2)

    (ICI / "participants").mkdir(exist_ok=True)
    ecrire_classeur(ICI / "Ventes_Retail.xlsx", ventes, FIN, v)
    ecrire_classeur(ICI / "Ventes_Retail_MAJ.xlsx", ventes_maj, FIN_MAJ, None)
    ecrire_classeur(ICI / "participants" / "Ventes_Retail.xlsx", ventes, FIN, None)
    ecrire_classeur(ICI / "participants" / "Ventes_Retail_MAJ.xlsx", ventes_maj, FIN_MAJ, None)

    v_maj = valeurs_attendues(ventes_maj, 0, 0)
    v["maj_lignes_apres_nettoyage"] = v_maj["lignes_apres_nettoyage"]
    v["maj_ca_total"] = v_maj["ca_total"]
    (ICI / "valeurs_attendues.json").write_text(json.dumps(v, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: x for k, x in v.items() if k != "mensuel"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
